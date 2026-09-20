from openpyxl import load_workbook

path = r'c:\privWork\workspace\webProject\docs\sample.xlsx'
wb = load_workbook(path)
for sheet_name in ['hrm_org_dept', 'hrm_person', 'hrm_assignment']:
    if sheet_name not in wb.sheetnames:
        continue
    ws = wb[sheet_name]
    print('SHEET', sheet_name, 'rows=', ws.max_row, 'cols=', ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(4, ws.max_row), values_only=True):
        print(row)
    print('---')
