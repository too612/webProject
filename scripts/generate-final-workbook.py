from pathlib import Path
import re
from datetime import date, timedelta
from openpyxl import Workbook, load_workbook

BASE = Path(r"c:\privWork\workspace\webProject")
TABLE_DIR = BASE / "docs" / "table"
OUT_PATH = BASE / "docs" / "sample.xlsx"

SHEET_TABLES = {
    "com_code": "com_code",
    "com_history": "com_history",
    "hrm_assignment": "hrm_assignment",
    "hrm_career": "hrm_career",
    "hrm_department": "hrm_org_dept",
    "hrm_person": "hrm_person",
    "hrm_worship_time": "hrm_worship_time",
}

SOURCE_COLUMNS = {
    "hrm_assignment": ["employee_no", "assignment_date", "assignment_type_code", "dept_cd", "employment_type_code", "grade_code", "position_code", "job_title_code", "assignment_content", "assignment_end_date", "concurrent_assignment_yn", "dispatch_country_code"],
    "hrm_career": ["employee_no", "hire_date", "retire_date", "company_name", "employment_type_code", "job_title", "job_responsibility"],
    "hrm_education": ["employee_no", "school_type_code", "admission_date", "graduation_date", "school_name", "graduation_status_code", "degree_code", "field_code", "major", "country_code", "is_final_education"],
    "hrm_person": ["employee_no", "name_ko", "name_en", "name_hanja", "dept_cd", "dept_name", "grade_code", "employment_type_code", "service_status_code", "birth_date", "gender_code", "hire_date", "retire_date", "promotion_date", "role_name", "ai_payload"],
}

LABELS = {
    "dept_cd": "부서코드",
    "dept_nm": "부서명",
    "eng_dept_nm": "부서영문명",
    "parent_dept_key": "상위부서키",
    "dept_level": "부서레벨",
    "change_type_cd": "변경구분코드",
    "chef_emp_no": "부서장사번",
    "dept_sys_cd": "원시스템코드",
    "remark": "비고",
    "employee_no": "사번",
    "name_ko": "이름",
    "name_en": "영문이름",
    "name_hanja": "한자이름",
    "grade_code": "직급코드",
    "position_code": "직위코드",
    "employment_type_code": "고용형태코드",
    "service_status_code": "재직구분코드",
    "birth_date": "생년월일",
    "gender_code": "성별코드",
    "hire_date": "입사일자",
    "retire_date": "퇴사일자",
    "promotion_date": "승진일자",
    "assignment_date": "발령일자",
    "assignment_type_code": "발령유형코드",
    "assignment_content": "발령내용",
    "assignment_end_date": "발령종료일",
    "concurrent_assignment_yn": "겸직여부",
    "dispatch_country_code": "파견국가코드",
    "year_label": "연도표시",
    "year_no": "연도",
    "event_date": "이벤트일",
    "description": "설명",
    "images": "이미지",
    "category": "분류",
    "title": "제목",
    "time": "시간",
    "location": "장소",
    "order_no": "순서",
    "country_code": "국가코드",
    "country_name": "국가명",
    "country_name_ko": "국가명(한글)",
    "country_name_en": "국가명(영문)",
    "code_group": "코드그룹",
    "code_value": "코드값",
    "code_name": "코드명",
    "sort_order": "정렬순서",
    "reorg_dt": "조직개편일자",
    "reorg_key": "조직개편키",
    "dept_key": "부서내부키",
    "person_key": "사원내부키",
    "assignment_key": "발령내부키",
    "history_id": "연혁ID",
    "reg_user": "등록자",
    "reg_dtm": "등록일시",
    "reg_ip": "등록IP",
    "upd_user": "수정자",
    "upd_dtm": "수정일시",
    "upd_ip": "수정IP",
    "ai_profile": "AI프로필",
    "extra_attributes": "추가속성",
    "header_key": "헤더키",
    "headline": "제목",
    "summary": "요약",
    "event_id": "이벤트ID",
    "use_yn": "사용여부",
    "conf_yn": "확정여부",
    "role_name": "역할명",
    "role_id": "역할ID",
}


def to_label(column_name: str) -> str:
    return LABELS.get(column_name, column_name.replace('_', ' '))


def split_top_level(text: str, delim: str = ','):
    parts = []
    buf = []
    depth = 0
    in_single = False
    for i, ch in enumerate(text):
        if ch == "'" and (i == 0 or text[i - 1] != "\\"):
            in_single = not in_single
        elif not in_single:
            if ch == '(':
                depth += 1
            elif ch == ')':
                depth -= 1
            elif ch == delim and depth == 0:
                part = ''.join(buf).strip()
                if part:
                    parts.append(part)
                buf = []
                continue
        buf.append(ch)
    tail = ''.join(buf).strip()
    if tail:
        parts.append(tail)
    return parts


def parse_literal(token: str):
    token = token.strip()
    if token in ('', 'NULL'):
        return ''
    token = re.sub(r'::[A-Za-z0-9_]+$', '', token)
    token = token.replace('DATE ', '').replace('TIMESTAMPTZ ', '').replace('INET ', '')
    if token.startswith("'") and token.endswith("'"):
        return token[1:-1].replace("''", "'")
    if token.startswith('"') and token.endswith('"'):
        return token[1:-1]
    if re.fullmatch(r'-?\d+', token):
        return int(token)
    if re.fullmatch(r'-?\d+\.\d+', token):
        return float(token)
    return token


def parse_ddl_columns(sql_text: str):
    matches = []
    pattern = re.compile(r'CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?([A-Za-z0-9_]+)\s*\(', re.I)
    for m in pattern.finditer(sql_text):
        table_name = m.group(1)
        close_idx = find_matching_parenthesis(sql_text, m.end() - 1)
        if close_idx is None:
            continue
        body = re.sub(r'--[^\r\n]*', '', sql_text[m.end():close_idx])
        cols = []
        for part in split_top_level(body):
            part = part.strip()
            if not part:
                continue
            if re.match(r'^(CONSTRAINT|PRIMARY|UNIQUE|FOREIGN|CHECK|INDEX|COMMENT|CREATE|ALTER|ON|USING|WITH|WITHOUT)\b', part, re.I):
                continue
            m2 = re.match(r'([A-Za-z0-9_]+)\b', part)
            if m2:
                cols.append(m2.group(1))
        matches.append((table_name, cols))
    return matches


def build_target_map():
    pairs = {}
    ddl_files = {p.stem: p for p in TABLE_DIR.glob('*.sql') if not p.name.endswith('_seed.sql')}
    seed_files = {p.stem[:-5]: p for p in TABLE_DIR.glob('*_seed.sql')}
    for sheet_name, table_name in SHEET_TABLES.items():
        ddl_path = ddl_files.get('hrm_department' if sheet_name == 'hrm_department' else table_name)
        seed_path = seed_files.get(sheet_name)
        if ddl_path is None or seed_path is None:
            continue
        ddl_text = ddl_path.read_text(encoding='utf-8')
        columns = []
        for name, candidate_cols in parse_ddl_columns(ddl_text):
            if name == table_name:
                columns = candidate_cols
                break
        if columns:
            pairs[sheet_name] = {'table': table_name, 'ddl': ddl_path, 'seed': seed_path, 'columns': columns}
    return pairs


def parse_seed_rows(seed_text: str, table_name: str = ''):
    rows = []
    segments = [seed_text]
    if table_name:
        insert_names = set(re.findall(r'\bINSERT\s+INTO\s+([A-Za-z0-9_]+)', seed_text, re.I))
        matches = list(re.finditer(rf'INSERT\s+INTO\s+{re.escape(table_name)}\b', seed_text, re.I))
        if matches and insert_names - {table_name}:
            segments = []
            for index, match in enumerate(matches):
                next_insert = re.search(r'\bINSERT\s+INTO\b', seed_text[match.end():], re.I)
                end = match.end() + next_insert.start() if next_insert else len(seed_text)
                segments.append(seed_text[match.start():end])

    for segment in segments:
        segment = re.sub(r'--[^\r\n]*', '', segment)
        for values_match in re.finditer(r'VALUES\s*\(', segment, re.I):
            cursor = values_match.end() - 1
            while cursor < len(segment) and segment[cursor] == '(':
                close_idx = find_matching_parenthesis(segment, cursor)
                if close_idx is None:
                    break
                tuple_text = segment[cursor + 1:close_idx]
                cells = [parse_literal(v) for v in split_top_level(tuple_text)]
                if cells and any(str(c) != '' for c in cells):
                    rows.append(cells)
                cursor = close_idx + 1
                while cursor < len(segment) and segment[cursor].isspace():
                    cursor += 1
                if cursor >= len(segment) or segment[cursor] != ',':
                    break
                cursor += 1
                while cursor < len(segment) and segment[cursor].isspace():
                    cursor += 1
    if not rows:
        rows = parse_union_select_rows(seed_text, table_name)
    return rows


def parse_insert_columns(seed_text: str, table_name: str):
    match = re.search(rf'INSERT\s+INTO\s+{re.escape(table_name)}\s*\((.*?)\)', seed_text, re.I | re.S)
    if not match:
        return []
    return [part.strip() for part in match.group(1).split(',') if part.strip()]


def extract_array_values(seed_text: str, marker: str):
    match = re.search(r'\(ARRAY\[(.*?)\]\)\[gs\].*?' + re.escape(marker), seed_text, re.S | re.I)
    if not match:
        return []
    return re.findall(r"'((?:''|[^'])*)'", match.group(1))


def expand_person_generated_rows(seed_text: str):
    rows = []
    arrays = [re.findall(r"'((?:''|[^'])*)'", match.group(1)) for match in re.finditer(r'\(ARRAY\[(.*?)\]\)\[gs\]', seed_text, re.S | re.I)]
    deacon_names = next((values for values in arrays if len(values) == 50), [])
    if len(deacon_names) >= 50:
        departments = [f'D{index:06d}' for index in range(9, 21)]
        for gs, name in enumerate(deacon_names[:50], 1):
            rows.append([
                f'{27 + gs:06d}', name, '', '', departments[(gs - 1) % 12], '',
                '103-070', '104-010' if gs % 3 == 0 else '104-020' if gs % 3 == 1 else '104-030',
                '101-010', date(1970, 1, 1) + timedelta(days=gs * 11),
                'M' if ((gs - 1) % 12) in (0, 1, 2, 3, 10) else 'F',
                date(2010, 1, 1) + timedelta(days=gs * 5), '', '', '집사', ''
            ])
    saint_names = next((values for values in reversed(arrays) if len(values) == 20), [])
    if len(saint_names) >= 20:
        for gs, name in enumerate(saint_names[:20], 1):
            rows.append([
                f'{77 + gs:06d}', name, '', '', 'D000019' if gs % 2 == 0 else 'D000020', '',
                '103-080', '104-030', '101-010', date(1992, 1, 1) + timedelta(days=gs * 17),
                'F' if gs % 2 == 0 else 'M', date(2018, 1, 1) + timedelta(days=gs * 3), '', '', '성도', ''
            ])
    return rows


def place_rows(columns, rows, input_columns):
    if not input_columns:
        return rows
    positions = {name: index for index, name in enumerate(columns)}
    placed = []
    for row in rows:
        target = [''] * len(columns)
        for index, value in enumerate(row[:len(input_columns)]):
            column_index = positions.get(input_columns[index])
            if column_index is not None:
                target[column_index] = value
        placed.append(target)
    return placed


def find_matching_parenthesis(text: str, open_idx: int):
    depth = 1
    in_single = False
    in_double = False
    for i in range(open_idx + 1, len(text)):
        ch = text[i]
        if ch == "'" and not in_double and (i == 0 or text[i - 1] != "\\"):
            in_single = not in_single
        elif ch == '"' and not in_single and (i == 0 or text[i - 1] != "\\"):
            in_double = not in_double
        elif not in_single and not in_double:
            if ch == '(':
                depth += 1
            elif ch == ')':
                depth -= 1
                if depth == 0:
                    return i
    return None


def parse_union_select_rows(seed_text: str, table_name: str = ''):
    rows = []
    for line in seed_text.splitlines():
        line = line.strip()
        if not re.match(r'(?:UNION\s+ALL\s+)?SELECT\b', line, re.I):
            continue
        line = re.sub(r'^(?:UNION\s+ALL\s+)?SELECT\s+', '', line, flags=re.I)
        line = re.sub(r'\s+FROM\s+.*$', '', line, flags=re.I)
        line = re.sub(r'\s+AS\s+\w+\s*$', '', line, flags=re.I)
        expressions = split_top_level(line)
        if not expressions:
            continue
        values = []
        for expression in expressions:
            expression = re.sub(r'\s+AS\s+\w+\s*$', '', expression, flags=re.I)
            values.append(parse_literal(expression))
        if values and not (len(values) == 1 and str(values[0]).upper() == 'SELECT'):
            rows.append(values)
    return rows


def add_sheet(wb: Workbook, table_name: str, columns, rows):
    ws = wb.create_sheet(title=table_name[:31])
    col_count = max(len(columns), max((len(r) for r in rows), default=0))
    logical = [to_label(c) for c in columns[:col_count]] + [""] * max(0, col_count - len(columns[:col_count]))
    physical = [c for c in columns[:col_count]] + [""] * max(0, col_count - len(columns[:col_count]))
    ws.append(logical)
    ws.append(physical)
    for row in rows:
        norm = list(row[:col_count]) + [""] * max(0, col_count - len(row[:col_count]))
        ws.append(norm)
    for col in ws.columns:
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 30)


def main():
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    targets = build_target_map()
    for table_key in sorted(targets):
        info = targets[table_key]
        ddl_path = info.get('ddl')
        seed_path = info.get('seed')
        if ddl_path is None or seed_path is None:
            continue
        ddl_text = ddl_path.read_text(encoding='utf-8')
        seed_text = seed_path.read_text(encoding='utf-8')
        columns = info.get('columns') or []
        if not columns:
            parsed = parse_ddl_columns(ddl_text)
            for name, candidate_cols in parsed:
                if name == table_key:
                    columns = candidate_cols
                    break
        if not columns:
            continue
        table_name = info.get('table', table_key)
        rows = parse_seed_rows(seed_text, table_name)
        if table_key == 'hrm_person':
            rows.extend(expand_person_generated_rows(seed_text))
        if table_key == 'com_code':
            additions = TABLE_DIR / 'com_code_seed_additions.sql'
            if additions.exists():
                rows.extend(parse_seed_rows(additions.read_text(encoding='utf-8'), table_name))
        if table_key == 'hrm_department' and not any(len(row) > 2 and row[2] == 'D000000' for row in rows):
            rows.insert(0, ['', '', 'D000000', '다사랑교회', 'Dasarang Church', '', 1, 'NEW', '000002', 'D000000', 'Y', '최상위 조직', 'SYSTEM', '', '127.0.0.1', 'SYSTEM', '', '127.0.0.1', '', ''])
        if not rows:
            continue
        input_columns = SOURCE_COLUMNS.get(table_key) or parse_insert_columns(seed_text, table_name)
        rows = place_rows(columns, rows, input_columns)
        add_sheet(wb, table_key, columns, rows)
    if len(wb.sheetnames) == 0:
        raise RuntimeError('No table data was generated into the workbook.')
    wb.save(OUT_PATH)
    print('saved=', OUT_PATH.exists())
    print('sheets=', wb.sheetnames)
    loaded = load_workbook(OUT_PATH)
    print('loaded_sheets=', loaded.sheetnames)
    first = loaded[loaded.sheetnames[0]]
    for row in first.iter_rows(min_row=1, max_row=min(5, first.max_row), values_only=True):
        print(row)


if __name__ == '__main__':
    main()
